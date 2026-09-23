#!/usr/bin/env python3
"""export the public interpreter snapshot from the private workbook"""

from __future__ import annotations

import json
import re
import sys
import unicodedata
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "data" / "AttestedLanguages_InterpretativeServices.xlsx"
OUTPUT = ROOT / "data" / "interpreters.json"
DETAIL_SHEET = "Interpreters Details"
LINK_SHEET = "Interpreters by Language"
ALIASES = {
    "Cal Interpreting and Translations": "Cal Interpreting & Translations (CIT)",
    "LanguageLine": "LanguageLine Solutions",
    "Languages - BIG Language Solutions": "BIG Language Solutions",
    "Languages | Maya Bridge": "Maya Bridge",
}
OMIT_WITHOUT_DETAIL = {"RIO: Indigenous Languages Fund"}


def clean(value: object) -> str | None:
    if value is None:
        return None
    value = str(value).strip()
    return value or None


def cell_url(cell, *, fallback_text: bool = True) -> str | None:
    """prefer hyperlink targets while accepting literal URLs"""
    target = cell.hyperlink.target if cell.hyperlink else None
    value = clean(target if target else (cell.value if fallback_text else None))
    if value and value.casefold() in {"n/a", "na", "none", "same as primary"}:
        return None
    return value


def slug(name: str) -> str:
    normalized = unicodedata.normalize("NFKD", name).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", "-", normalized).strip("-")


def require(condition: bool, message: str) -> None:
    if not condition:
        raise ValueError(message)


def build_snapshot(source: Path) -> dict:
    # read-only openpyxl cells dont expose hyperlink targets
    workbook = load_workbook(source, data_only=True, read_only=False)
    require(DETAIL_SHEET in workbook.sheetnames, f"Missing sheet: {DETAIL_SHEET}")
    require(LINK_SHEET in workbook.sheetnames, f"Missing sheet: {LINK_SHEET}")
    details = workbook[DETAIL_SHEET]
    header = {clean(cell.value): cell.column for cell in details[1] if clean(cell.value)}

    def detail_value(row: int, field: str):
        return details.cell(row, header[field])

    required_fields = {
        "service_name", "service_type", "oregon_relevance", "languages_published",
        "primary_source_url", "Contact/Request URL", "phone_number", "contact_email",
    }
    normalized_header = {key.rstrip(): column for key, column in header.items()}
    # contact header has a trailing space in this workbook
    header = normalized_header
    missing = required_fields - header.keys()
    require(not missing, f"Missing provider columns: {', '.join(sorted(missing))}")

    providers = []
    names_to_ids: dict[str, str] = {}
    ids = set()
    for row in range(2, details.max_row + 1):
        name = clean(detail_value(row, "service_name").value)
        if not name:
            continue
        provider_id = slug(name)
        require(provider_id and provider_id not in ids, f"Provider ID collision: {name!r} -> {provider_id!r}")
        ids.add(provider_id)
        primary = cell_url(detail_value(row, "primary_source_url"))
        request_cell = detail_value(row, "Contact/Request URL")
        request_text = clean(request_cell.value)
        if request_text and request_text.casefold() == "same as primary":
            request = primary
        else:
            request = cell_url(request_cell)
        provider = {
            "id": provider_id,
            "name": name,
            "serviceType": clean(detail_value(row, "service_type").value),
            "oregonRelevance": clean(detail_value(row, "oregon_relevance").value),
            "languagesPublished": clean(detail_value(row, "languages_published").value),
            "primarySourceUrl": primary,
            "requestUrl": request,
            "phoneNumber": clean(detail_value(row, "phone_number").value),
            "contactEmail": clean(detail_value(row, "contact_email").value),
        }
        for url_field in ("primarySourceUrl", "requestUrl"):
            value = provider[url_field]
            require(value is None or is_web_url(value), f"Invalid {url_field} for {name}: {value!r}")
        providers.append(provider)
        names_to_ids[name] = provider_id

    links_sheet = workbook[LINK_SHEET]
    link_headers = {clean(cell.value): cell.column for cell in links_sheet[1] if clean(cell.value)}
    require({"Language", "Provider", "Value"} <= link_headers.keys(), "Unexpected Interpreters by Language headers")
    language_providers = set()
    unresolved = set()
    for row in range(2, links_sheet.max_row + 1):
        language = clean(links_sheet.cell(row, link_headers["Language"]).value)
        provider_name = clean(links_sheet.cell(row, link_headers["Provider"]).value)
        value = clean(links_sheet.cell(row, link_headers["Value"]).value)
        if not language or not provider_name or (value or "").casefold() != "yes":
            continue
        canonical_name = ALIASES.get(provider_name, provider_name)
        provider_id = names_to_ids.get(canonical_name)
        if provider_id is None:
            if provider_name in OMIT_WITHOUT_DETAIL:
                continue
            unresolved.add(provider_name)
            continue
        language_providers.add((language, provider_id))

    require(not unresolved, "Unresolved pivot providers: " + ", ".join(sorted(unresolved)))
    linked_ids = {provider_id for _, provider_id in language_providers}
    providers = [provider for provider in providers if provider["id"] in linked_ids]
    providers.sort(key=lambda item: item["id"])
    links = [{"language": lang, "providerId": pid} for lang, pid in sorted(language_providers, key=lambda x: (x[0].casefold(), x[0], x[1]))]
    snapshot = {"providers": providers, "languageProviders": links}
    validate_snapshot(snapshot)
    return snapshot


def is_web_url(value: str) -> bool:
    parsed = urlparse(value)
    return parsed.scheme in {"http", "https"} and bool(parsed.netloc)


def validate_snapshot(snapshot: dict) -> None:
    require(set(snapshot) == {"providers", "languageProviders"}, "Unexpected top-level JSON fields")
    providers = snapshot["providers"]
    links = snapshot["languageProviders"]
    fields = {"id", "name", "serviceType", "oregonRelevance", "languagesPublished", "primarySourceUrl", "requestUrl", "phoneNumber", "contactEmail"}
    ids = [p.get("id") for p in providers]
    require(len(ids) == len(set(ids)), "Provider IDs are not unique")
    require(all(set(p) == fields for p in providers), "Provider record has missing or private extra fields")
    for p in providers:
        require(p["id"] == slug(p["name"]), f"Non-deterministic provider ID for {p['name']}")
        require(bool(p["name"]), "Provider name must not be empty")
        require(p["primarySourceUrl"] is None or is_web_url(p["primarySourceUrl"]), f"Bad primary URL for {p['name']}")
        require(p["requestUrl"] is None or is_web_url(p["requestUrl"]), f"Bad request URL for {p['name']}")
    seen = set()
    id_set = set(ids)
    for link in links:
        require(set(link) == {"language", "providerId"}, "Unexpected language link fields")
        require(bool(link["language"]), "Language name must not be empty")
        require(link["providerId"] in id_set, f"Unknown provider ID: {link['providerId']}")
        pair = (link["language"], link["providerId"])
        require(pair not in seen, f"Duplicate language-provider link: {pair}")
        seen.add(pair)


def main() -> int:
    try:
        snapshot = build_snapshot(SOURCE)
        OUTPUT.parent.mkdir(parents=True, exist_ok=True)
        OUTPUT.write_text(json.dumps(snapshot, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    except (OSError, KeyError, ValueError) as exc:
        print(f"Interpreter export failed: {exc}", file=sys.stderr)
        return 1
    print(f"Wrote {len(snapshot['providers'])} providers and {len(snapshot['languageProviders'])} language links to {OUTPUT.relative_to(ROOT)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
